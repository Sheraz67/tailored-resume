#!/usr/bin/env python3
"""Resume Tailor Web App — Flask backend."""

import io
import json
import os
import re
import tempfile
from datetime import datetime
from urllib.parse import urlparse

import requests
from flask import Flask, render_template, request, jsonify, send_file
from fpdf import FPDF
import anthropic
from google import genai
from google.genai import types

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 10 * 1024 * 1024  # 10 MB max upload


# ─────────────────────────────────────────────
# Resume file parsing
# ─────────────────────────────────────────────

def parse_docx(file_bytes: bytes) -> str:
    from docx import Document
    doc = Document(io.BytesIO(file_bytes))
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())


def parse_pdf(file_bytes: bytes) -> str:
    from PyPDF2 import PdfReader
    reader = PdfReader(io.BytesIO(file_bytes))
    text_parts = []
    for page in reader.pages:
        t = page.extract_text()
        if t:
            text_parts.append(t)
    return "\n".join(text_parts)


def parse_txt(file_bytes: bytes) -> str:
    return file_bytes.decode("utf-8", errors="replace")


PARSERS = {
    ".docx": parse_docx,
    ".pdf": parse_pdf,
    ".txt": parse_txt,
}


# ─────────────────────────────────────────────
# Default tailoring prompt
# ─────────────────────────────────────────────

DEFAULT_TAILORING_PROMPT = """
================================================================================
MASTER PROMPT: RESUME TAILORING ENGINE (v1.0)
================================================================================

You are a world-class Resume Strategist and ATS Optimization Expert with 15+ years
of experience in technical recruiting, hiring pipeline optimization, and career
coaching for senior software engineers. You have deep expertise in how Applicant
Tracking Systems (Greenhouse, Lever, Workday, iCIMS, Taleo) parse, score, and rank
resumes. You understand exactly how hiring managers, technical recruiters, and
engineering directors evaluate resumes for senior/staff-level engineering roles.

================================================================================
YOUR MISSION
================================================================================

You will receive TWO inputs:
  1. A CANDIDATE RESUME (the source of truth for experience, skills, and facts)
  2. A JOB DESCRIPTION (the target you are optimizing toward)

Your task is to produce a TAILORED RESUME that maximizes the candidate's chances of:
  (a) Passing ATS keyword filters and scoring algorithms
  (b) Getting shortlisted by a technical recruiter (6-second scan test)
  (c) Impressing the hiring manager during a detailed read
  (d) Setting up strong talking points for interviews

================================================================================
PHASE 1: DEEP JD ANALYSIS (Do this FIRST before touching the resume)
================================================================================

Before making ANY changes to the resume, perform this comprehensive JD analysis.
Output your analysis so the reasoning is transparent.

1.1 — ROLE IDENTITY EXTRACTION
    - What is the EXACT title? (Use this to retitle the resume)
    - What is the seniority level? (Junior / Mid / Senior / Staff / Principal)
    - What is the role TYPE? (IC, Tech Lead, Architect, Manager-of-one, Player-Coach)
    - Is it frontend-heavy, backend-heavy, full-stack equal, or has a PRIMARY focus?
    - What DOMAIN is the company in? (Healthcare, Finance, E-commerce, SaaS, etc.)
    - Remote / Hybrid / On-site?

1.2 — HARD REQUIREMENTS EXTRACTION (Must-Have)
    Extract EVERY hard requirement. Categorize them:

    [LANGUAGES]       — e.g., TypeScript, C#, Python
    [FRAMEWORKS]      — e.g., React, Angular, .NET, Next.js
    [ARCHITECTURE]    — e.g., Microservices, Event-Driven, Distributed Systems
    [CLOUD/DEVOPS]    — e.g., AWS, Azure, Docker, Kubernetes, CI/CD
    [DATA]            — e.g., PostgreSQL, MongoDB, SQL Server
    [AI/ML]           — e.g., LLM, RAG, Inference Pipelines, Model Deployment
    [SECURITY]        — e.g., Secure Coding, DevSecOps, Compliance
    [SOFT SKILLS]     — e.g., Mentoring, Cross-functional Collaboration
    [YEARS]           — Minimum years of experience required
    [EDUCATION]       — Degree requirements if any

1.3 — PREFERRED/BONUS QUALIFICATIONS
    Extract ALL preferred qualifications separately. These are differentiators.
    Mark each as [HIGH-VALUE BONUS] or [NICE-TO-HAVE].

1.4 — HIDDEN KEYWORDS & IMPLICIT REQUIREMENTS
    Identify keywords that are IMPLIED but not explicitly stated:
    - If JD says "scalable frontend architectures" -> implies: Component Libraries,
      Design Systems, State Management, Performance Budgets
    - If JD says "production ownership" -> implies: On-call, Incident Response,
      Runbooks, SLA Management, Post-mortems
    - If JD says "AI integration" -> implies: Prompt Engineering, Token Optimization,
      Model Evaluation, A/B Testing of AI features
    - If JD says "observability" -> implies: DataDog/Grafana/NewRelic, Distributed
      Tracing, Log Aggregation, Alerting Thresholds

1.5 — TONE & CULTURE SIGNALS
    Analyze the JD's language for cultural signals:
    - Does it say "ownership" / "accountability" -> Emphasize autonomous leadership
    - Does it say "collaborate" / "cross-functional" -> Emphasize teamwork
    - Does it say "systems thinking" -> Emphasize architectural judgment
    - Does it say "measurable impact" -> Every bullet MUST have metrics
    - Does it say "not a feature-delivery role" -> De-emphasize "built features",
      emphasize "owned systems", "designed architecture", "improved reliability"

1.6 — PRIORITY STACK RANKING
    Rank all extracted requirements by importance (based on JD ordering, repetition,
    and emphasis). The resume should mirror this priority order:

    Priority 1 (MUST dominate resume): _______________
    Priority 2 (MUST be prominent):    _______________
    Priority 3 (Should appear clearly): _______________
    Priority 4 (Should be mentioned):   _______________
    Priority 5 (Nice to include):       _______________

================================================================================
PHASE 2: RESUME AUDIT (Map existing resume against JD)
================================================================================

2.1 — KEYWORD MATCH MATRIX
    Create a matrix:
    | JD Requirement          | Present in Resume? | Where?          | Strength |
    |-------------------------|--------------------|-----------------|----------|
    | TypeScript              | YES                | Summary, Skills | Strong   |
    | Core Web Vitals         | NO                 | -               | Missing  |
    | Event-Driven APIs       | Partial            | Microsoft only  | Weak     |

    Mark each as: STRONG / ADEQUATE / WEAK / MISSING

2.2 — EXPERIENCE ALIGNMENT SCORING
    For each work experience entry, score alignment to JD (0-100%):
    - Which bullets already align?
    - Which bullets are irrelevant to this JD?
    - Which bullets can be REFRAMED to align?
    - What MISSING experiences from this role should be surfaced?

2.3 — GAP IDENTIFICATION
    List ALL gaps between JD and resume:
    [CRITICAL GAPS]   — Required skills/keywords completely missing
    [LANGUAGE GAPS]    — Skill exists but wrong terminology used
    [EMPHASIS GAPS]    — Skill exists but buried or underweighted
    [STRUCTURAL GAPS]  — Information in wrong section or wrong order

================================================================================
PHASE 3: RESUME REWRITING RULES
================================================================================

FOLLOW THESE RULES WITH ZERO EXCEPTIONS:

3.1 — GOLDEN RULE: NEVER FABRICATE
    - NEVER invent experience, skills, projects, metrics, or companies
    - NEVER inflate numbers beyond what's stated in the original resume
    - NEVER add technologies the candidate hasn't used
    - You may REFRAME, REWORD, and REORDER — but NEVER FABRICATE
    - If a JD requirement has NO match in the resume, leave it out — do NOT fake it

3.2 — TITLE OPTIMIZATION
    - Match the resume title to the JD title as closely as possible
    - If JD says "Senior Frontend-Focused Full Stack Engineer", the resume title
      should reflect frontend-first positioning
    - Do NOT copy the JD title verbatim if it sounds unnatural — adapt it

3.3 — SUMMARY REWRITING (Most critical section — recruiters read this first)

    Structure (4-6 sentences max):
    Sentence 1: [Seniority] + [Primary Identity from JD] + [Years] + [Domain Match]
    Sentence 2: [PRIMARY technical strength matching JD Priority 1]
    Sentence 3: [SECONDARY technical strength matching JD Priority 2]
    Sentence 4: [AI/Cloud/Specialized skill matching JD Priority 3]
    Sentence 5: [Ownership/Leadership/Impact signal matching JD tone]
    Sentence 6: [Differentiator — regulated industry, scale, mentoring]

    Rules:
    - First 15 words must contain the JD's top 2-3 keywords
    - Include EXACT phrases from JD where natural (ATS matching)
    - Quantify at least 2 claims (years, scale, percentage improvements)
    - Mirror the JD's tone (ownership-focused vs. collaboration-focused)
    - Do NOT use first person ("I built...") — use telegraphic style
    - Do NOT use buzzwords without substance ("passionate", "rockstar", "guru")

3.4 — SKILLS SECTION REWRITING

    Rules:
    - REORDER skill categories to match JD priority stack (Phase 1.6)
    - RENAME skill categories to mirror JD section headers EXACTLY
      (If JD says "Frontend Architecture & Performance Leadership",
       the skill category should be "Frontend Architecture & Performance")
    - Within each category, list JD-mentioned skills FIRST, then supporting skills
    - ADD skills that exist in the resume but are scattered in experience bullets
    - REMOVE or DEPRIORITIZE skills not relevant to this specific JD
    - Use the EXACT terminology from the JD:
      JD says "C#" -> use "C#" (not "C Sharp")
      JD says "Kubernetes" -> use "Kubernetes" (not "K8s")
      JD says "event-driven" -> use "Event-Driven" (not "message-based")
    - Include IMPLICIT skills from Phase 1.4 if candidate has them

3.5 — EXPERIENCE BULLET REWRITING (The most impactful section)

    For EACH bullet point, apply this framework:

    FORMULA: [POWER VERB] + [WHAT you did using JD KEYWORDS] + [SCALE/SCOPE] +
             [MEASURABLE RESULT with METRIC]

    Example transformation:
    BEFORE: "Built React components for dashboards"
    AFTER:  "Architected scalable React and TypeScript frontend systems for clinical
             dashboards serving 10k+ daily users, optimizing Core Web Vitals and
             reducing Time-to-Interactive by 40%"

    Power Verb Selection (match to JD tone):
    - For architecture roles: Architected, Designed, Engineered, Established
    - For ownership roles:    Owned, Led, Drove, Spearheaded
    - For optimization roles: Optimized, Reduced, Improved, Accelerated
    - For AI/innovation:      Implemented, Integrated, Developed, Productionized
    - For mentoring:          Mentored, Guided, Established, Influenced

    AVOID these weak verbs: Helped, Assisted, Participated, Worked on, Was responsible

    Rules for each bullet:
    a) Start with the STRONGEST action verb (no two bullets should start with same verb)
    b) Front-load JD keywords in the first half of the bullet
    c) Include at least ONE metric per bullet (%, #, scale, time saved)
    d) Connect the action to BUSINESS IMPACT (not just technical output)
    e) Use JD's EXACT phrases where natural
    f) If a bullet doesn't serve this specific JD, either REFRAME it or REMOVE it
    g) Aim for 5-7 bullets per recent role, 4-5 for older roles
    h) Most impactful/JD-aligned bullets go FIRST within each role

3.6 — EXPERIENCE ORDERING WITHIN EACH ROLE
    Reorder bullets within each job to match JD priority:
    1st bullet: PRIMARY JD focus (e.g., Frontend Architecture)
    2nd bullet: SECONDARY JD focus (e.g., Backend/Systems)
    3rd bullet: AI/ML integration (if applicable)
    4th bullet: Cloud/DevOps/Observability
    5th bullet: Data/Security
    6th bullet: Leadership/Mentoring
    7th bullet: Other measurable impact

3.7 — COMPANY CONTEXT ANNOTATIONS
    For each role, add a brief context tag if the company domain matches JD preferences:
    - If JD prefers regulated industries -> tag "(Healthcare - Regulated Industry)"
    - If JD prefers high-scale -> tag "(E-Commerce - 2M+ Monthly Users)"
    - This signals domain alignment immediately to the recruiter

3.8 — EDUCATION & CERTIFICATIONS
    - Keep education section brief
    - If candidate has relevant certifications (AWS, Azure, etc.), ensure they're visible
    - If JD mentions specific certifications, and candidate has them, HIGHLIGHT them

================================================================================
PHASE 4: ATS OPTIMIZATION CHECKLIST
================================================================================

After rewriting, verify ALL of the following:

4.1 — KEYWORD DENSITY
    [ ] Every REQUIRED skill from JD appears at least 2x in resume
        (once in Skills, once in Experience)
    [ ] Every PREFERRED skill appears at least 1x
    [ ] Job title or close variant appears in resume title AND summary
    [ ] Industry-specific terms from JD are present

4.2 — FORMATTING (ATS-Safe)
    [ ] No tables, columns, graphics, or images (ATS can't parse them)
    [ ] No headers/footers with critical info (ATS often skips them)
    [ ] Standard section headers: SUMMARY, SKILLS, EXPERIENCE, EDUCATION
    [ ] Dates in consistent format (MM/YYYY or Month YYYY)
    [ ] No special characters that might break parsing (em-dashes, smart quotes)
    [ ] Plain text or simple formatting only
    [ ] Contact info at the very top (not in a sidebar)

4.3 — LENGTH & DENSITY
    [ ] 2-3 pages maximum for 10+ years experience
    [ ] No single bullet exceeds 3 lines
    [ ] White space is adequate for readability
    [ ] Most recent 2 roles get the most space (60% of experience section)

4.4 — CONSISTENCY
    [ ] Same tense throughout (past tense for previous roles, present for current)
    [ ] Consistent bullet formatting (all start with verb, all have metrics)
    [ ] Technology names spelled identically throughout (TypeScript, not TS sometimes)
    [ ] No orphan skills (skills in Skills section should appear in Experience too)

================================================================================
PHASE 5: FINAL QUALITY CHECKS
================================================================================

Before delivering, verify:

5.1 — THE 6-SECOND TEST
    Read ONLY the title, summary, and skill headers. In 6 seconds, is it clear that
    this candidate matches the JD? If not, rewrite.

5.2 — THE KEYWORD MATCH TEST
    Highlight every JD keyword in the tailored resume. Coverage should be 85%+.
    List any JD keywords NOT present and explain why they were excluded.

5.3 — THE AUTHENTICITY TEST
    Compare every claim in the tailored resume against the original resume.
    Flag ANYTHING that could be seen as exaggerated or fabricated.

5.4 — THE "SO WHAT" TEST
    Every bullet should answer: "So what? Why does this matter to THIS employer?"
    If a bullet doesn't connect to JD priorities, cut it or reframe it.

5.5 — THE SPECIFICITY TEST
    Remove any vague language:
    - "various technologies" -> name the specific technologies
    - "improved performance" -> "reduced API latency by 25%"
    - "worked with teams" -> "mentored 5 engineers across 3 teams"
    - "large-scale systems" -> "systems handling 5M+ records/month"

5.6 — THE RECENCY BIAS CHECK
    Most recent role should have the MOST bullets and STRONGEST alignment.
    Older roles can be condensed — but don't lose critical JD-matching experience.

================================================================================
CONSTRAINTS & ANTI-PATTERNS
================================================================================

NEVER DO THESE:
  x  Do NOT copy-paste JD language into resume without adapting it naturally
  x  Do NOT use the same metric in multiple bullets
  x  Do NOT start consecutive bullets with the same word
  x  Do NOT include technologies the candidate hasn't used (even if JD requires them)
  x  Do NOT remove ALL non-JD experience (shows breadth and avoids looking "too perfect")
  x  Do NOT use subjective self-assessments ("excellent communicator", "strong leader")
  x  Do NOT include objective statements ("Seeking a role in...")
  x  Do NOT include references or "References available upon request"
  x  Do NOT use acronyms without spelling them out at least once
  x  Do NOT create a resume longer than 3 pages
  x  Do NOT sacrifice readability for keyword stuffing

================================================================================
NOW EXECUTE
================================================================================

Apply ALL phases above to the following inputs:
""".strip()


# ─────────────────────────────────────────────
# AI API calls (Claude + Gemini)
# ─────────────────────────────────────────────

def _build_tailor_user_msg(resume_text: str, prompt_text: str, jd_text: str) -> str:
    """Build the full prompt for resume tailoring (shared across providers)."""
    system_msg = prompt_text.strip()
    user_msg = f"""Here are the two inputs:

=== CANDIDATE RESUME ===
{resume_text}

=== JOB DESCRIPTION ===
{jd_text}

=== INSTRUCTIONS ===
Apply every phase from your system prompt to these inputs.

IMPORTANT: Return ONLY the tailored resume as a JSON object with this exact structure (no change log, no interview prep — just the resume):

```json
{{
  "name": "Candidate Name",
  "title": "Tailored Job Title",
  "contact": "Location | email",
  "summary": "The full summary paragraph",
  "skills": [
    {{"category": "Category Name", "items": "skill1, skill2, skill3"}}
  ],
  "experience": [
    {{
      "job_title": "Title",
      "company": "Company Name",
      "context": "Domain context tag",
      "dates": "MM/YYYY - MM/YYYY",
      "location": "City, ST or Remote",
      "bullets": ["bullet 1", "bullet 2"]
    }}
  ],
  "education": {{
    "degree": "Degree Name",
    "school": "School Name",
    "dates": "MM/YYYY - MM/YYYY",
    "location": "City, ST"
  }}
}}
```

Return ONLY the JSON object, no markdown code fences, no extra text. Just pure JSON.
"""
    return f"{system_msg}\n\n{user_msg}"


def _strip_code_fences(raw: str) -> str:
    """Strip markdown code fences from AI response."""
    raw = raw.strip()
    if raw.startswith("```"):
        raw = re.sub(r"^```(?:json)?\s*", "", raw)
        raw = re.sub(r"\s*```$", "", raw)
    return raw


def _call_claude(api_key: str, prompt: str, max_tokens: int = 8000) -> str:
    """Call Claude and return the raw text response."""
    client = anthropic.Anthropic(api_key=api_key)
    message = client.messages.create(
        model="claude-sonnet-4-5-20250929",
        max_tokens=max_tokens,
        messages=[{"role": "user", "content": prompt}],
    )
    return message.content[0].text.strip()


def _call_gemini(api_key: str, prompt: str, max_tokens: int = 8000) -> str:
    """Call Gemini and return the raw text response."""
    client = genai.Client(api_key=api_key)
    response = client.models.generate_content(
        model="gemini-2.0-flash",
        contents=prompt,
        config=types.GenerateContentConfig(
            max_output_tokens=max_tokens,
        ),
    )
    return response.text.strip()


def call_ai(provider: str, api_key: str, prompt: str, max_tokens: int = 8000) -> str:
    """Route to the correct AI provider and return raw text."""
    if provider == "gemini":
        return _call_gemini(api_key, prompt, max_tokens)
    return _call_claude(api_key, prompt, max_tokens)


def tailor_resume(api_key: str, resume_text: str, prompt_text: str, jd_text: str, provider: str = "claude") -> dict:
    """Call AI to tailor the resume. Returns structured JSON."""
    prompt = _build_tailor_user_msg(resume_text, prompt_text, jd_text)
    raw = call_ai(provider, api_key, prompt)
    raw = _strip_code_fences(raw)
    return json.loads(raw)


# ─────────────────────────────────────────────
# PDF generation
# ─────────────────────────────────────────────

def _clean_text(text: str) -> str:
    """Replace Unicode characters unsupported by Helvetica with ASCII equivalents."""
    return (text
        .replace("\u2013", "-")   # en dash
        .replace("\u2014", "-")   # em dash
        .replace("\u2018", "'")   # left single quote
        .replace("\u2019", "'")   # right single quote
        .replace("\u201c", '"')   # left double quote
        .replace("\u201d", '"')   # right double quote
        .replace("\u2022", "-")   # bullet
        .replace("\u2026", "...")  # ellipsis
        .replace("\u00a0", " ")   # non-breaking space
        .replace("\u2010", "-")   # hyphen
        .replace("\u2011", "-")   # non-breaking hyphen
        .replace("\u2012", "-")   # figure dash
        .replace("\u2015", "-")   # horizontal bar
        .replace("\u2027", "-")   # hyphenation point
        .replace("\u2032", "'")   # prime
        .replace("\u2033", '"')   # double prime
    )


class ResumePDF(FPDF):
    BLUE = (44, 95, 138)
    DARK = (26, 26, 26)
    BODY = (51, 51, 51)
    GRAY = (102, 102, 102)

    def __init__(self):
        super().__init__()
        self.set_auto_page_break(auto=True, margin=15)

    def header(self):
        pass  # no header

    def section_header(self, text):
        self.set_font("Helvetica", "B", 12)
        self.set_text_color(*self.BLUE)
        self.cell(0, 8, _clean_text(text).upper(), new_x="LMARGIN", new_y="NEXT")
        self.set_draw_color(*self.BLUE)
        self.line(self.l_margin, self.get_y(), self.w - self.r_margin, self.get_y())
        self.ln(3)

    def body_text(self, text, size=10):
        self.set_font("Helvetica", "", size)
        self.set_text_color(*self.BODY)
        self.multi_cell(0, 5, _clean_text(text))

    def bullet(self, text):
        x = self.get_x()
        self.set_font("Helvetica", "", 9.5)
        self.set_text_color(*self.BODY)
        indent = 8
        self.set_x(x + indent)
        self.multi_cell(self.w - self.r_margin - self.get_x(), 4.5, f"- {_clean_text(text)}")
        self.ln(0.5)


def generate_pdf(data: dict) -> bytes:
    pdf = ResumePDF()
    pdf.add_page()
    pdf.set_margins(18, 15, 18)
    pdf.set_y(15)

    # Name
    pdf.set_font("Helvetica", "B", 20)
    pdf.set_text_color(*ResumePDF.DARK)
    pdf.cell(0, 10, _clean_text(data.get("name", "")), align="C", new_x="LMARGIN", new_y="NEXT")

    # Title
    pdf.set_font("Helvetica", "", 13)
    pdf.set_text_color(*ResumePDF.BLUE)
    pdf.cell(0, 7, _clean_text(data.get("title", "")), align="C", new_x="LMARGIN", new_y="NEXT")

    # Contact
    pdf.set_font("Helvetica", "", 9.5)
    pdf.set_text_color(*ResumePDF.GRAY)
    pdf.cell(0, 6, _clean_text(data.get("contact", "")), align="C", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(4)

    # Summary
    pdf.section_header("Summary")
    pdf.body_text(data.get("summary", ""))
    pdf.ln(3)

    # Skills
    pdf.section_header("Skills")
    for skill in data.get("skills", []):
        cat = _clean_text(skill.get("category", ""))
        items = _clean_text(skill.get("items", ""))
        pdf.set_font("Helvetica", "B", 9.5)
        pdf.set_text_color(*ResumePDF.BODY)
        label = f"{cat}: "
        pdf.cell(pdf.get_string_width(label) + 1, 5, label)
        pdf.set_font("Helvetica", "", 9.5)
        remaining = pdf.w - pdf.r_margin - pdf.get_x()
        pdf.multi_cell(remaining, 4.5, items)
        pdf.ln(1)
    pdf.ln(2)

    # Experience
    pdf.section_header("Experience")
    for job in data.get("experience", []):
        pdf.set_font("Helvetica", "B", 11)
        pdf.set_text_color(*ResumePDF.DARK)
        pdf.cell(0, 6, _clean_text(job.get("job_title", "")), new_x="LMARGIN", new_y="NEXT")

        pdf.set_font("Helvetica", "", 9.5)
        pdf.set_text_color(*ResumePDF.BLUE)
        company_ctx = _clean_text(f"{job.get('company', '')} -- {job.get('context', '')}")
        pdf.cell(pdf.get_string_width(company_ctx) + 2, 5, company_ctx)
        pdf.set_text_color(*ResumePDF.GRAY)
        meta = _clean_text(f"  |  {job.get('dates', '')}  |  {job.get('location', '')}")
        pdf.cell(0, 5, meta, new_x="LMARGIN", new_y="NEXT")
        pdf.ln(1)

        for b in job.get("bullets", []):
            pdf.bullet(b)
        pdf.ln(3)

    # Education
    edu = data.get("education", {})
    if edu:
        pdf.section_header("Education")
        pdf.set_font("Helvetica", "B", 10.5)
        pdf.set_text_color(*ResumePDF.BODY)
        pdf.cell(0, 6, _clean_text(edu.get("degree", "")), new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Helvetica", "", 10)
        pdf.set_text_color(*ResumePDF.GRAY)
        edu_meta = _clean_text(f"{edu.get('school', '')}  |  {edu.get('dates', '')}  |  {edu.get('location', '')}")
        pdf.cell(0, 5, edu_meta)

    return pdf.output()


# ─────────────────────────────────────────────
# JD scraping helpers
# ─────────────────────────────────────────────

PLATFORM_DOMAINS = {
    "linkedin.com": "LinkedIn",
    "indeed.com": "Indeed",
    "glassdoor.com": "Glassdoor",
    "ziprecruiter.com": "ZipRecruiter",
    "monster.com": "Monster",
    "dice.com": "Dice",
    "lever.co": "Lever",
    "greenhouse.io": "Greenhouse",
    "workday.com": "Workday",
    "myworkdayjobs.com": "Workday",
    "smartrecruiters.com": "SmartRecruiters",
    "angel.co": "AngelList",
    "wellfound.com": "Wellfound",
    "builtin.com": "Built In",
    "simplyhired.com": "SimplyHired",
    "careerbuilder.com": "CareerBuilder",
    "welcometothejungle.com": "Welcome to the Jungle",
}

# All known platform names (lowercase) for filtering from titles
_PLATFORM_NAMES_LOWER = {v.lower() for v in PLATFORM_DOMAINS.values()} | {
    "linkedin", "indeed", "indeed.com", "glassdoor", "glassdoor.com",
    "otta", "welcome to the jungle",
}


def _is_platform_part(text):
    """Check if a title part is a platform name (not a real company)."""
    t = text.lower().strip()
    for pn in _PLATFORM_NAMES_LOWER:
        if pn in t:
            return True
    return False


def extract_job_metadata(url, title, text):
    """Detect platform from domain and parse company/position from page title."""
    parsed = urlparse(url)
    domain = parsed.netloc.lower().replace("www.", "")

    # Detect platform
    platform = "Other"
    for domain_key, platform_name in PLATFORM_DOMAINS.items():
        if domain_key in domain:
            platform = platform_name
            break

    company = ""
    position = ""

    if title:
        # LinkedIn: "Job Title at Company | LinkedIn"
        # Indeed: "Job Title - Company - Location | Indeed.com"
        # Glassdoor: "Company hiring Job Title in Location | Glassdoor"
        # WTTJ: "Company - Role | Welcome to the Jungle"

        if platform == "LinkedIn":
            m = re.match(r"^(.+?)\s+at\s+(.+?)(?:\s*\||\s*[-–]|\s*$)", title)
            if m:
                position = m.group(1).strip()
                company = m.group(2).strip()
        elif platform == "Indeed":
            parts = re.split(r"\s*[-–]\s*", title)
            if len(parts) >= 2:
                position = parts[0].strip()
                company = parts[1].strip()
        elif platform == "Glassdoor":
            m = re.match(r"^(.+?)\s+hiring\s+(.+?)(?:\s+in\s+|\s*\|)", title)
            if m:
                company = m.group(1).strip()
                position = m.group(2).strip()
        elif platform == "Welcome to the Jungle":
            # "Company - Role | Welcome to the Jungle (formerly Otta)"
            # Split on | first, take everything before the platform name
            main = re.split(r"\s*\|\s*", title)[0]
            parts = re.split(r"\s*[-–]\s*", main, maxsplit=1)
            if len(parts) >= 2:
                company = parts[0].strip()
                position = parts[1].strip()
            elif len(parts) == 1:
                position = parts[0].strip()

        # Generic fallback
        if not company and not position:
            # Split on | first to separate "content | platform branding"
            pipe_parts = re.split(r"\s*\|\s*", title)
            # Filter out parts that are platform names
            content_parts = [p.strip() for p in pipe_parts if not _is_platform_part(p)]
            if not content_parts:
                content_parts = [p.strip() for p in pipe_parts]

            # Use just the first content part (before any |), split on dash for company/role
            main = content_parts[0]
            sub = re.split(r"\s*[-–]\s*", main, maxsplit=1)
            if len(sub) >= 2:
                company = sub[0].strip()
                position = sub[1].strip()
            elif len(sub) == 1:
                position = sub[0].strip()

    return {
        "platform": platform,
        "company": company,
        "position": position,
        "url": url,
    }


# ─────────────────────────────────────────────
# Routes
# ─────────────────────────────────────────────

@app.route("/")
def index():
    return render_template("index.html")


@app.route("/api/scrape-jd", methods=["POST"])
def api_scrape_jd():
    """Scrape a job description from a URL using Apify website-content-crawler."""
    try:
        data = request.get_json()
        if not data:
            return jsonify({"error": "No data provided."}), 400

        apify_token = data.get("apify_token", "").strip()
        url = data.get("url", "").strip()

        if not apify_token:
            return jsonify({"error": "Apify API token is required."}), 400
        if not url:
            return jsonify({"error": "Job URL is required."}), 400

        # Call Apify website-content-crawler actor
        actor_id = "apify~website-content-crawler"
        run_url = f"https://api.apify.com/v2/acts/{actor_id}/runs"

        run_input = {
            "startUrls": [{"url": url}],
            "maxCrawlPages": 1,
            "crawlerType": "playwright:firefox",
            "maxConcurrency": 1,
            "proxyConfiguration": {"useApifyProxy": True},
        }

        # Start the actor run synchronously (wait for finish)
        resp = requests.post(
            run_url,
            params={"token": apify_token, "waitForFinish": 120},
            json=run_input,
            timeout=180,
        )

        if resp.status_code == 401:
            return jsonify({"error": "Invalid Apify API token."}), 401
        if not resp.ok:
            return jsonify({"error": f"Apify API error: {resp.status_code} {resp.text[:200]}"}), 502

        run_data = resp.json().get("data", {})
        run_status = run_data.get("status")

        if run_status != "SUCCEEDED":
            return jsonify({"error": f"Apify crawl did not succeed (status: {run_status}). Try again."}), 502

        # Fetch results from the default dataset
        dataset_id = run_data.get("defaultDatasetId")
        if not dataset_id:
            return jsonify({"error": "No dataset returned from Apify."}), 502

        items_url = f"https://api.apify.com/v2/datasets/{dataset_id}/items"
        items_resp = requests.get(
            items_url,
            params={"token": apify_token, "format": "json"},
            timeout=30,
        )

        if not items_resp.ok:
            return jsonify({"error": "Failed to fetch crawl results."}), 502

        items = items_resp.json()
        if not items:
            return jsonify({"error": "No content was extracted from the URL. The page may require login."}), 404

        item = items[0]
        page_text = item.get("text", "")
        page_title = item.get("metadata", {}).get("title", "") or item.get("title", "")

        if not page_text:
            return jsonify({"error": "Page was loaded but no text content was found."}), 404

        metadata = extract_job_metadata(url, page_title, page_text)

        return jsonify({
            "success": True,
            "text": page_text,
            "title": page_title,
            "metadata": metadata,
        })

    except requests.Timeout:
        return jsonify({"error": "Apify request timed out. Try again."}), 504
    except Exception as e:
        return jsonify({"error": f"Scraping failed: {str(e)}"}), 500


@app.route("/api/tailor", methods=["POST"])
def api_tailor():
    try:
        provider = request.form.get("provider", "claude").strip().lower()
        api_key = request.form.get("api_key", "").strip()
        if not api_key:
            label = "Gemini" if provider == "gemini" else "Anthropic"
            return jsonify({"error": f"{label} API key is required."}), 400

        jd_text = request.form.get("jd", "").strip()
        resume_text_direct = request.form.get("resume_text", "").strip()

        # Parse prompt from file or text
        prompt_text = request.form.get("prompt", "").strip()
        if "prompt_file" in request.files:
            pf = request.files["prompt_file"]
            if pf.filename:
                pext = os.path.splitext(pf.filename)[1].lower()
                if pext not in PARSERS:
                    return jsonify({"error": f"Unsupported prompt file type: {pext}. Use .txt, .docx, or .pdf"}), 400
                prompt_text = PARSERS[pext](pf.read())

        if not prompt_text:
            prompt_text = DEFAULT_TAILORING_PROMPT
        if not jd_text:
            return jsonify({"error": "Job description is required."}), 400

        # Parse resume from file or text
        resume_text = ""
        if "resume_file" in request.files:
            f = request.files["resume_file"]
            if f.filename:
                ext = os.path.splitext(f.filename)[1].lower()
                if ext not in PARSERS:
                    return jsonify({"error": f"Unsupported file type: {ext}. Use .docx, .pdf, or .txt"}), 400
                resume_text = PARSERS[ext](f.read())

        if not resume_text and resume_text_direct:
            resume_text = resume_text_direct

        if not resume_text:
            return jsonify({"error": "Please upload a resume file or paste resume text."}), 400

        result = tailor_resume(api_key, resume_text, prompt_text, jd_text, provider)
        return jsonify({"success": True, "data": result})

    except json.JSONDecodeError as e:
        return jsonify({"error": f"Failed to parse AI response as JSON. Try again. Details: {str(e)}"}), 500
    except anthropic.AuthenticationError:
        return jsonify({"error": "Invalid API key. Please check your Anthropic API key."}), 401
    except anthropic.RateLimitError:
        return jsonify({"error": "Rate limited. Please wait a moment and try again."}), 429
    except Exception as e:
        err_str = str(e)
        if "API_KEY_INVALID" in err_str or "401" in err_str:
            return jsonify({"error": "Invalid API key. Please check your Gemini API key."}), 401
        return jsonify({"error": f"Something went wrong: {err_str}"}), 500


@app.route("/api/answer-questions", methods=["POST"])
def api_answer_questions():
    """Answer job application questions using JD + resume context."""
    try:
        data = request.get_json()
        if not data:
            return jsonify({"error": "No data provided."}), 400

        provider = data.get("provider", "claude").strip().lower()
        api_key = data.get("api_key", "").strip()
        questions = data.get("questions", "").strip()
        jd_text = data.get("jd", "").strip()
        resume_json = data.get("resume", {})

        if not api_key:
            return jsonify({"error": "API key is required."}), 400
        if not questions:
            return jsonify({"error": "Please paste at least one question."}), 400
        if not jd_text:
            return jsonify({"error": "Job description context is missing."}), 400

        # Rebuild a readable resume summary from the structured data
        resume_lines = []
        resume_lines.append(f"Name: {resume_json.get('name', '')}")
        resume_lines.append(f"Title: {resume_json.get('title', '')}")
        resume_lines.append(f"\nSummary:\n{resume_json.get('summary', '')}")
        resume_lines.append("\nSkills:")
        for s in resume_json.get("skills", []):
            resume_lines.append(f"  {s.get('category', '')}: {s.get('items', '')}")
        resume_lines.append("\nExperience:")
        for job in resume_json.get("experience", []):
            resume_lines.append(f"\n  {job.get('job_title', '')} at {job.get('company', '')} ({job.get('dates', '')})")
            for b in job.get("bullets", []):
                resume_lines.append(f"    - {b}")
        edu = resume_json.get("education", {})
        if edu:
            resume_lines.append(f"\nEducation: {edu.get('degree', '')} — {edu.get('school', '')}")
        resume_text = "\n".join(resume_lines)

        system_prompt = """You are an expert job application assistant. You help candidates write compelling,
authentic answers to job application questions. You have deep context about:
1. The candidate's background (their tailored resume)
2. The specific job they are applying for (the job description)

Rules:
- Write answers in FIRST PERSON as the candidate
- Keep answers concise but substantive (3-6 sentences per question unless it clearly needs more)
- Ground every answer in REAL experience from the resume — never fabricate
- Mirror the tone and keywords from the job description naturally
- Show enthusiasm for the specific role and company
- If a question asks about something not covered in the resume, craft an honest answer
  that pivots to relevant strengths rather than making things up
- For salary questions, suggest the candidate research market rates rather than giving a number
- For "why this company" questions, reference specific things from the JD that align with the candidate's experience"""

        user_msg = f"""Here is the candidate's resume:

{resume_text}

Here is the job description they are applying to:

{jd_text}

Please answer each of the following application questions. Format your response as a JSON array where each
element has "question" (the original question) and "answer" (your crafted response).

Questions:
{questions}

Return ONLY a JSON array, no markdown code fences, no extra text. Example format:
[{{"question": "Why do you want this role?", "answer": "Your answer here..."}}]"""

        full_prompt = f"{system_prompt}\n\n{user_msg}"
        raw = call_ai(provider, api_key, full_prompt, max_tokens=4000)
        raw = _strip_code_fences(raw)

        answers = json.loads(raw)
        return jsonify({"success": True, "answers": answers})

    except json.JSONDecodeError as e:
        return jsonify({"error": f"Failed to parse AI response. Try again. Details: {str(e)}"}), 500
    except anthropic.AuthenticationError:
        return jsonify({"error": "Invalid API key."}), 401
    except anthropic.RateLimitError:
        return jsonify({"error": "Rate limited. Please wait and try again."}), 429
    except Exception as e:
        err_str = str(e)
        if "API_KEY_INVALID" in err_str or "401" in err_str:
            return jsonify({"error": "Invalid API key. Please check your Gemini API key."}), 401
        return jsonify({"error": f"Something went wrong: {err_str}"}), 500


@app.route("/api/download-pdf", methods=["POST"])
def api_download_pdf():
    try:
        data = request.get_json()
        if not data:
            return jsonify({"error": "No resume data provided."}), 400

        pdf_bytes = generate_pdf(data)
        name_slug = data.get("name", "resume").replace(" ", "_")
        filename = f"{name_slug}.pdf"

        return send_file(
            io.BytesIO(pdf_bytes),
            mimetype="application/pdf",
            as_attachment=True,
            download_name=filename,
        )
    except Exception as e:
        return jsonify({"error": f"PDF generation failed: {str(e)}"}), 500


if __name__ == "__main__":
    app.run(debug=True, port=5000)
